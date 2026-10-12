"""Nạp HÓA ĐƠN CŨ (2017 – 03/2026) từ file Excel KiotViet cũ vào anhkhach.db để ghép ảnh.

KiotViet hiện tại chỉ có hóa đơn từ 04/2026. Hóa đơn trước đó nằm trên Google Drive, thư mục
"HÓA ĐƠN 2017- QUÝ 1 2026" (HD-2017.xlsx … HD-2025.xlsx, HD_ QUY1-2026.xlsx) — là bản xuất
"Danh sách chi tiết hóa đơn" của KiotViet cũ: mỗi dòng 1 món hàng, nhiều dòng cùng Mã hóa đơn.

    .venv/bin/python hoadon_cu.py ~/hoadon_cu/*.xlsx            # chạy khan: chỉ đếm, không ghi
    .venv/bin/python hoadon_cu.py ~/hoadon_cu/*.xlsx --ghi      # ghi vào anhkhach.db

Mã hóa đơn cũ có thể trùng mã của KiotViet mới (cùng dạng HD0xxxxx) -> lưu với tiền tố "CU-".
Chỉ đọc file Excel, không sửa file trên Drive. Hóa đơn "Đã hủy" bị bỏ.
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import sys

from anhkhach import DB_FILE, chuan_sdt, ket_noi

TIEN_TO = "CU-"
COT = {  # tên cột trong file xuất KiotViet -> khóa dùng ở đây
    "Mã hóa đơn": "ma", "Thời gian": "tg", "Mã khách hàng": "ma_kh", "Tên khách hàng": "ten",
    "Điện thoại": "sdt", "Tên hàng": "hang", "Chi nhánh": "cn", "Trạng thái": "tt",
}


def _ngay_gio(v) -> tuple[str, str] | None:
    if isinstance(v, dt.datetime):
        return v.date().isoformat(), v.strftime("%H:%M")
    s = str(v or "").strip()
    for f in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            t = dt.datetime.strptime(s, f)
            return t.date().isoformat(), t.strftime("%H:%M")
        except ValueError:
            pass
    return None


def doc_file(path: str) -> dict[str, dict]:
    """Gộp các dòng món hàng thành hóa đơn: {mã: {ngay, gio, ma_kh, ten, sdt, dv[], cn}}."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    hd: dict[str, dict] = {}
    for ws in wb.worksheets:
        vi_tri = None
        for hang in ws.iter_rows(values_only=True):
            if vi_tri is None:  # tìm dòng tiêu đề (có "Mã hóa đơn")
                ten_cot = [str(x or "").strip() for x in hang]
                if "Mã hóa đơn" in ten_cot:
                    vi_tri = {COT[c]: i for i, c in enumerate(ten_cot) if c in COT}
                continue

            def lay(k):
                i = vi_tri.get(k)
                return hang[i] if i is not None and i < len(hang) else None

            ma = str(lay("ma") or "").strip()
            ng = _ngay_gio(lay("tg"))
            if not ma or not ng or "hủy" in str(lay("tt") or "").lower():
                continue
            h = hd.setdefault(ma, {"ngay": ng[0], "gio": ng[1], "ma_kh": str(lay("ma_kh") or "").strip(),
                                   "ten": str(lay("ten") or "").strip() or "Khách lẻ",
                                   "sdt": chuan_sdt(str(lay("sdt") or "")), "dv": [],
                                   "cn": str(lay("cn") or "").strip()})
            ten_hang = str(lay("hang") or "").strip()
            if ten_hang and ten_hang not in h["dv"]:
                h["dv"].append(ten_hang)
    wb.close()
    return hd


def main() -> None:
    ap = argparse.ArgumentParser(description="Nạp hóa đơn cũ (Excel KiotViet cũ) vào anhkhach.db")
    ap.add_argument("files", nargs="+")
    ap.add_argument("--ghi", action="store_true", help="ghi thật (mặc định chỉ đếm)")
    a = ap.parse_args()
    tat_ca: dict[str, dict] = {}
    for f in a.files:
        hd = doc_file(os.path.expanduser(f))
        ngays = sorted(h["ngay"] for h in hd.values())
        print(f"{os.path.basename(f)}: {len(hd)} hóa đơn"
              + (f" ({ngays[0]} → {ngays[-1]})" if ngays else " — không thấy cột 'Mã hóa đơn'"))
        tat_ca.update(hd)
    theo_thang: dict[str, int] = {}
    for h in tat_ca.values():
        theo_thang[h["ngay"][:7]] = theo_thang.get(h["ngay"][:7], 0) + 1
    print("Theo tháng:", ", ".join(f"{k}: {v}" for k, v in sorted(theo_thang.items())[-15:]))
    if not a.ghi:
        print("CHẠY KHAN — chưa ghi. Thêm --ghi để ghi vào", os.path.abspath(DB_FILE))
        return
    conn = ket_noi()
    # Mã khách KiotViet cũ có thể là mã của NGƯỜI KHÁC bên KiotViet mới -> khớp theo SĐT trước;
    # không khớp thì giữ mã cũ với tiền tố CU- (thà tách 2 hồ sơ còn hơn gộp nhầm 2 người).
    theo_sdt: dict[str, set] = {}
    for sdt, ma_kh in conn.execute(f"SELECT sdt, ma_kh FROM visits WHERE code NOT LIKE '{TIEN_TO}%'"):
        if sdt and ma_kh:
            theo_sdt.setdefault(sdt, set()).add(ma_kh)
    khop = 0
    for ma, h in tat_ca.items():
        moi = theo_sdt.get(h["sdt"]) if h["sdt"] else None
        if moi and len(moi) == 1:
            ma_kh, khop = next(iter(moi)), khop + 1
        else:
            ma_kh = TIEN_TO + h["ma_kh"] if h["ma_kh"] else ""
        conn.execute("INSERT OR REPLACE INTO visits VALUES (?,?,?,?,?,?,?,?)",
                     (TIEN_TO + ma, h["ngay"], h["gio"], ma_kh, h["ten"], h["sdt"],
                      " | ".join(h["dv"]), h["cn"]))
    print(f"  {khop} hóa đơn cũ nối được với mã khách KiotViet mới qua SĐT")
    conn.commit()
    print(f"Đã ghi {len(tat_ca)} hóa đơn cũ (mã {TIEN_TO}…). Giờ chạy: anhkhach.py ghep YYYY-MM")


if __name__ == "__main__":
    sys.exit(main())
